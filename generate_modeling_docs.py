from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches

from src.config import Config


def paragraph_block(markdown_text: str, docx_text: str | None = None) -> dict[str, str]:
    return {
        "type": "paragraph",
        "md": markdown_text,
        "docx": docx_text if docx_text is not None else markdown_text,
    }


def bullet_block(markdown_text: str, docx_text: str | None = None) -> dict[str, str]:
    return {
        "type": "bullet",
        "md": markdown_text,
        "docx": docx_text if docx_text is not None else markdown_text,
    }


def equation_block(
    label: str,
    markdown_formula: str,
    docx_formula: str,
    image_formula=None,
) -> dict[str, str]:
    return {
        "type": "equation",
        "label": label,
        "md": markdown_formula,
        "docx": docx_formula,
        "image": image_formula if image_formula is not None else markdown_formula,
    }


def build_sections() -> list[tuple[str, list[dict[str, str]]]]:
    return [
        (
            "1. 模型改进动机",
            [
                paragraph_block(
                    "原始模型将充电时间简化为 `充电量 / 恒定充电功率`，并默认车辆到站后可立即开始充电，因此无法刻画高 SOC 区间充电变慢以及热门站点排队等待的现象。",
                    "原始模型将充电时间简化为“充电量 / 恒定充电功率”，并默认车辆到站后可立即开始充电，因此无法刻画高 SOC 区间充电变慢以及热门站点排队等待的现象。",
                ),
                paragraph_block(
                    "原始模型还将路段运行速度近似为按时间段切换的全局常数，没有显式链路、没有容量约束，也无法体现多车辆同时上路所引起的拥塞传播与回溢。",
                ),
                paragraph_block(
                    "为提升模型的物理真实性，本文在保留既有成本结构的基础上，引入了基于 SOC 的非线性充电模型、共享充电桩排队模型，以及带有限存储能力的链路时变拥塞网络。",
                ),
            ],
        ),
        (
            "2. 符号与集合",
            [
                bullet_block(
                    "车辆集合记为 $V$，充电站集合记为 $S$，客户集合记为 $C$，合成路网的有向链路集合记为 $A$。",
                    "车辆集合记为 V，充电站集合记为 S，客户集合记为 C，合成路网的有向链路集合记为 A。",
                ),
                bullet_block(
                    f"对任意车辆 $v \\in V$，其电池容量为 $B={Config.BATTERY_CAPACITY:.1f}\\,\\text{{kWh}}$，载重容量为 $Q={Config.CAPACITY:.1f}$。",
                    f"对任意车辆 v∈V，其电池容量为 B={Config.BATTERY_CAPACITY:.1f} kWh，载重容量为 Q={Config.CAPACITY:.1f}。",
                ),
                bullet_block(
                    "车辆在时刻 $t$ 的荷电状态记为 $SOC_v(t) \\in [0,1]$，对应剩余电量为 $E_v(t)=B\\,SOC_v(t)$。",
                    "车辆在时刻 t 的荷电状态记为 SOC_v(t)∈[0,1]，对应剩余电量为 E_v(t)=B·SOC_v(t)。",
                ),
                bullet_block(
                    "对任意链路 $(i,j) \\in A$，其长度记为 $d_{ij}$，自由流速度记为 $u^0_{ij}$，基础通行能力记为 $q^0_{ij}$，存储能力记为 $K_{ij}$。",
                    "对任意链路 (i,j)∈A，其长度记为 d_ij，自由流速度记为 u_ij^0，基础通行能力记为 q_ij^0，存储能力记为 K_ij。",
                ),
                bullet_block(
                    "对任意充电站 $s \\in S$，其充电桩数量记为 $m_s$，单桩额定功率记为 $P_s^{\\max}$。",
                    "对任意充电站 s∈S，其充电桩数量记为 m_s，单桩额定功率记为 P_s^max。",
                ),
            ],
        ),
        (
            "3. 非线性充电模型",
            [
                paragraph_block(
                    "本文采用分段 SOC 充电曲线来描述电池在不同荷电状态下的充电功率变化。设站点额定功率为 $P_s^{\\max}$，则单位时刻有效充电功率表示为",
                    "本文采用分段 SOC 充电曲线来描述电池在不同荷电状态下的充电功率变化。设站点额定功率为 P_s^max，则单位时刻有效充电功率表示为",
                ),
                equation_block(
                    "(4-1)",
                    r"P_s(SOC)=P_s^{\max}\cdot \phi(SOC)",
                    "P_s(SOC) = P_s^max · φ(SOC)",
                ),
                paragraph_block(
                    "其中，$\\phi(SOC)$ 取分段函数：",
                    "其中，φ(SOC) 取分段函数：",
                ),
                equation_block(
                    "(4-2)",
                    r"\phi(SOC)=\begin{cases}1.0, & 0\le SOC<0.60 \\ 1.0-2.0(SOC-0.60), & 0.60\le SOC<0.85 \\ 0.5-2.0(SOC-0.85), & 0.85\le SOC\le 1.00\end{cases}",
                    "φ(SOC) = {1.0, 0≤SOC<0.60; 1.0 - 2.0(SOC - 0.60), 0.60≤SOC<0.85; 0.5 - 2.0(SOC - 0.85), 0.85≤SOC≤1.00}",
                    image_formula={
                        "kind": "cases",
                        "lhs": r"\phi(SOC)=",
                        "lines": [
                            r"1.0,\quad 0\leq SOC<0.60",
                            r"1.0-2.0(SOC-0.60),\quad 0.60\leq SOC<0.85",
                            r"0.5-2.0(SOC-0.85),\quad 0.85\leq SOC\leq 1.00",
                        ],
                    },
                ),
                paragraph_block(
                    "相应地，从 $SOC_a$ 充至 $SOC_b$ 的充电时长通过分段积分计算：",
                    "相应地，从 SOC_a 充至 SOC_b 的充电时长通过分段积分计算：",
                ),
                equation_block(
                    "(4-3)",
                    r"T_s^{chg}(SOC_a,SOC_b)=60\int_{SOC_a}^{SOC_b} \frac{B}{P_s(SOC)}\,dSOC",
                    "T_s^chg(SOC_a, SOC_b) = 60 ∫[SOC_a,SOC_b] B / P_s(SOC) dSOC",
                ),
                paragraph_block(
                    "在程序实现中，上式通过分段数值积分求解，从而避免线性充电假设对高 SOC 区间时长的系统性低估。",
                ),
                paragraph_block(
                    f"本文默认仓库、私有站、公共站的额定功率分别设为 {Config.DEPOT_CHARGING_POWER:.0f} kW、{Config.PRIVATE_CHARGING_POWER:.0f} kW 和 {Config.PUBLIC_CHARGING_POWER:.0f} kW。",
                ),
            ],
        ),
        (
            "4. 共享充电桩排队模型",
            [
                paragraph_block(
                    "考虑到同一时段内多辆车可能竞争同一站点有限的充电桩资源，本文采用有限服务台排队机制描述站内等待过程。",
                ),
                paragraph_block(
                    "对于充电站 $s$，设其共有 $m_s$ 个充电桩，每个充电桩维护一个“最早可用时刻”。当车辆 $v$ 在时刻 $t_v^{arr,s}$ 到达站点后，其开始充电时刻定义为",
                    "对于充电站 s，设其共有 m_s 个充电桩，每个充电桩维护一个“最早可用时刻”。当车辆 v 在时刻 t_v^{arr,s} 到达站点后，其开始充电时刻定义为",
                ),
                equation_block(
                    "(4-4)",
                    r"t_v^{start,s}=\max\left(t_v^{arr,s},\min_{k=1,\dots,m_s} a_{sk}\right)",
                    "t_v^{start,s} = max(t_v^{arr,s}, min_{k=1,...,m_s} a_sk)",
                ),
                paragraph_block(
                    "其中 $a_{sk}$ 为第 $k$ 个充电桩当前的最早可用时刻。由此得到排队等待时间",
                    "其中 a_sk 为第 k 个充电桩当前的最早可用时刻。由此得到排队等待时间",
                ),
                equation_block(
                    "(4-5)",
                    r"W_v^s=t_v^{start,s}-t_v^{arr,s}",
                    "W_v^s = t_v^{start,s} - t_v^{arr,s}",
                ),
                paragraph_block(
                    "车辆完成充电后的离站时刻为",
                ),
                equation_block(
                    "(4-6)",
                    r"t_v^{dep,s}=t_v^{start,s}+T_s^{chg}(SOC_a,SOC_b)",
                    "t_v^{dep,s} = t_v^{start,s} + T_s^chg(SOC_a, SOC_b)",
                ),
                paragraph_block(
                    "完成服务后，将对应充电桩的最早可用时刻更新为 $t_v^{dep,s}$。该机制使得不同车辆在同一解内共享站点服务状态，排队等待会进一步传导到时间窗罚金、制冷成本与后续路网占用中。",
                    "完成服务后，将对应充电桩的最早可用时刻更新为 t_v^{dep,s}。该机制使得不同车辆在同一解内共享站点服务状态，排队等待会进一步传导到时间窗罚金、制冷成本与后续路网占用中。",
                ),
                paragraph_block(
                    f"本文默认仓库、私有站、公共站的充电桩数量分别为 {Config.DEPOT_CHARGER_COUNT}、{Config.PRIVATE_CHARGER_COUNT}、{Config.PUBLIC_CHARGER_COUNT}。",
                ),
            ],
        ),
        (
            "5. 合成路网构造方法",
            [
                paragraph_block(
                    "由于原始数据仅提供仓库、客户与充电站坐标，缺乏真实道路拓扑，本文基于几何邻近关系自动生成合成有向路网。",
                ),
                paragraph_block(
                    f"具体地，对所有节点构造 $k$ 近邻图，本文取 $k={Config.NETWORK_K_NEIGHBORS}$；若初始图不连通，则逐步连接不同连通分量之间的最近节点对，直至整图连通。",
                    f"具体地，对所有节点构造 k 近邻图，本文取 k={Config.NETWORK_K_NEIGHBORS}；若初始图不连通，则逐步连接不同连通分量之间的最近节点对，直至整图连通。",
                ),
                paragraph_block(
                    "对每条无向边，依据其长度与全网边长中位数的比较，将其划分为“支路”或“干路”，并将其拆分为两条方向相反的有向链路。",
                ),
                paragraph_block(
                    f"支路与干路的自由流速度分别设置为 {Config.NETWORK_LOCAL_FREE_SPEED_KMH:.0f} km/h 和 {Config.NETWORK_ARTERIAL_FREE_SPEED_KMH:.0f} km/h，对应的基础通行能力分别为 {Config.NETWORK_LOCAL_BASE_CAPACITY_VPH:.0f} veh/h 和 {Config.NETWORK_ARTERIAL_BASE_CAPACITY_VPH:.0f} veh/h。",
                ),
            ],
        ),
        (
            "6. 链路时变旅行时间与拥塞传播",
            [
                paragraph_block(
                    "本文不再采用全局常速假设，而是在链路层面刻画时变运行时间。对任意链路 $(i,j)$，其峰时自由流速度与峰时通行能力分别按比例衰减：",
                    "本文不再采用全局常速假设，而是在链路层面刻画时变运行时间。对任意链路 (i,j)，其峰时自由流速度与峰时通行能力分别按比例衰减：",
                ),
                equation_block(
                    "(4-7)",
                    rf"u_{{ij}}(t)=u_{{ij}}^0\cdot {Config.NETWORK_PEAK_SPEED_FACTOR:.2f}\text{{，若 }} t \text{{ 落入峰时；否则 }} u_{{ij}}(t)=u_{{ij}}^0",
                    f"u_ij(t) = u_ij^0 · {Config.NETWORK_PEAK_SPEED_FACTOR:.2f}，若 t 落入峰时；否则 u_ij(t) = u_ij^0",
                    image_formula=[
                        rf"u_{{ij}}(t)=u_{{ij}}^0\cdot {Config.NETWORK_PEAK_SPEED_FACTOR:.2f},\quad t\in \mathrm{{peak}}",
                        r"u_{ij}(t)=u_{ij}^0,\quad \mathrm{otherwise}",
                    ],
                ),
                equation_block(
                    "(4-8)",
                    rf"q_{{ij}}(t)=q_{{ij}}^0\cdot {Config.NETWORK_PEAK_CAPACITY_FACTOR:.2f}\text{{，若 }} t \text{{ 落入峰时；否则 }} q_{{ij}}(t)=q_{{ij}}^0",
                    f"q_ij(t) = q_ij^0 · {Config.NETWORK_PEAK_CAPACITY_FACTOR:.2f}，若 t 落入峰时；否则 q_ij(t) = q_ij^0",
                    image_formula=[
                        rf"q_{{ij}}(t)=q_{{ij}}^0\cdot {Config.NETWORK_PEAK_CAPACITY_FACTOR:.2f},\quad t\in \mathrm{{peak}}",
                        r"q_{ij}(t)=q_{ij}^0,\quad \mathrm{otherwise}",
                    ],
                ),
                paragraph_block(
                    "链路自由流运行时间为",
                ),
                equation_block(
                    "(4-9)",
                    r"\tau_{ij}^{ff}(t)=\dfrac{60d_{ij}}{u_{ij}(t)}",
                    "τ_ij^ff(t) = 60 d_ij / u_ij(t)",
                ),
                paragraph_block(
                    "同时，为反映车辆积压与回溢，本文为每条链路设置有限存储能力",
                ),
                equation_block(
                    "(4-10)",
                    rf"K_{{ij}}=\rho^{{jam}}d_{{ij}}\text{{，其中 }} \rho^{{jam}}={Config.NETWORK_JAM_DENSITY_VEH_PER_KM:.0f}\,\text{{veh/km}}",
                    f"K_ij = ρ^jam d_ij，其中 ρ^jam = {Config.NETWORK_JAM_DENSITY_VEH_PER_KM:.0f} veh/km",
                    image_formula=rf"K_{{ij}}=\rho^{{jam}}d_{{ij}},\quad \rho^{{jam}}={Config.NETWORK_JAM_DENSITY_VEH_PER_KM:.0f}\ \mathrm{{veh/km}}",
                ),
                paragraph_block(
                    "若车辆准备进入链路时，该链路当前在网车辆数已达到 $K_{ij}$，则车辆必须在上游等待至最早释放时刻，形成回溢阻塞。",
                    "若车辆准备进入链路时，该链路当前在网车辆数已达到 K_ij，则车辆必须在上游等待至最早释放时刻，形成回溢阻塞。",
                ),
                paragraph_block(
                    "在出口侧，链路以服务率 $q_{ij}(t)$ 释放车辆，相邻车辆的最小出流时间间隔为",
                    "在出口侧，链路以服务率 q_ij(t) 释放车辆，相邻车辆的最小出流时间间隔为",
                ),
                equation_block(
                    "(4-11)",
                    r"h_{ij}(t)=\dfrac{60}{q_{ij}(t)}",
                    "h_ij(t) = 60 / q_ij(t)",
                ),
                paragraph_block(
                    "因此，车辆在链路上的实际离开时刻取自由流离开时刻与容量约束离开时刻中的较大者。该设计等价于带有限存储能力的点队列模型，既能体现时间依赖旅行时间，又能描述多车相互作用下的拥塞传播。",
                ),
            ],
        ),
        (
            "7. 与配送成本的耦合",
            [
                paragraph_block(
                    "车辆在任意一步决策中，需要同时比较“直接前往目标节点”和“先前往某充电站补能再前往目标节点”两类方案。",
                ),
                paragraph_block(
                    "对于任意候选动作，其总增量成本仍由运输距离成本、充电费用、时间窗罚金、货损成本和制冷成本组成，只是这些量不再由静态直线距离直接给出，而是由动态路网仿真与排队仿真共同决定。",
                ),
                paragraph_block(
                    "特别地，排队等待时间和拥塞延误时间不会作为独立罚金重复计费，而是通过延长在途时间和停留时间，间接影响时间窗罚金、制冷成本以及后续车辆的链路与站点占用状态。",
                ),
            ],
        ),
        (
            "8. 模型假设与说明",
            [
                paragraph_block(
                    "1. 由于缺乏真实 GIS 路网，本文使用节点坐标自动构造合成路网，重点在于体现链路层面的时变性与拥塞传播，而非复现实地道路几何细节。",
                ),
                paragraph_block(
                    "2. 链路拥塞采用有限存储点队列近似，不显式模拟车道变换与跟驰行为，但足以支持算法层面的时间依赖路径代价评估。",
                ),
                paragraph_block(
                    "3. 充电站等待时间采用共享充电桩调度计算，不再假定车辆到站即可立即开始充电，因此更适合描述公共充电资源竞争。",
                ),
                paragraph_block(
                    "4. 本文保留原有目标函数结构，使改进模型能够在不改变算法主框架输入输出形式的前提下，直接嵌入现有求解流程。",
                ),
            ],
        ),
    ]


def build_markdown_content() -> str:
    lines = ["# IGWO 模型补充说明", ""]
    for title, blocks in build_sections():
        lines.append(f"## {title}")
        lines.append("")
        for block in blocks:
            if block["type"] == "bullet":
                lines.append(f"- {block['md']}")
            elif block["type"] == "equation":
                lines.append(f"{block['label']}  ${block['md']}$")
            else:
                lines.append(block["md"])
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def set_run_font(run, western_font: str, east_asia_font: str | None = None) -> None:
    east_font = east_asia_font or western_font
    run.font.name = western_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_font)


def _parse_markup_token(text: str, start: int) -> tuple[str, int] | None:
    if start >= len(text) or text[start] not in "_^":
        return None

    marker = text[start]
    if start + 1 >= len(text):
        return None

    if text[start + 1] == "{":
        depth = 1
        cursor = start + 2
        while cursor < len(text) and depth > 0:
            if text[cursor] == "{":
                depth += 1
            elif text[cursor] == "}":
                depth -= 1
            cursor += 1
        if depth != 0:
            return None
        return marker + text[start + 2 : cursor - 1], cursor

    cursor = start + 1
    while cursor < len(text) and text[cursor] not in " \t\n\r()[]{}=+-/*,，；：。_^":
        cursor += 1
    if cursor == start + 1:
        return None
    return marker + text[start + 1 : cursor], cursor


def add_markup_runs(
    paragraph,
    text: str,
    western_font: str,
    east_asia_font: str,
    *,
    math_mode: bool = False,
) -> None:
    cursor = 0
    while cursor < len(text):
        token = _parse_markup_token(text, cursor)
        if token is not None:
            content, next_cursor = token
            run = paragraph.add_run(content[1:])
            set_run_font(run, western_font, east_asia_font)
            if content[0] == "_":
                run.font.subscript = True
            else:
                run.font.superscript = True
            cursor = next_cursor
            continue

        next_marker = len(text)
        for marker in ("_{", "^{", "_", "^"):
            found = text.find(marker, cursor)
            if found != -1:
                next_marker = min(next_marker, found)

        if next_marker == cursor:
            run = paragraph.add_run(text[cursor])
            set_run_font(run, western_font, east_asia_font)
            cursor += 1
            continue

        chunk = text[cursor:next_marker]
        if chunk:
            run = paragraph.add_run(chunk)
            set_run_font(run, western_font, east_asia_font)
        cursor = next_marker


def add_code_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, "Consolas", "Consolas")


def sanitize_formula_label(label: str) -> str:
    label_text = label.strip().strip("()")
    return "eq_" + label_text.replace("-", "_").replace(" ", "")


def _render_formula_lines(lines: list[str], output_path: Path) -> Path:
    max_line_length = max(len(line) for line in lines)
    fig_width = max(4.8, min(10.0, max_line_length * 0.12))
    fig_height = 0.65 * len(lines) + 0.35

    fig = plt.figure(figsize=(fig_width, fig_height))
    fig.patch.set_alpha(0.0)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis("off")

    y_positions = [0.75 - index * 0.28 for index in range(len(lines))]
    for y_pos, line in zip(y_positions, lines):
        formula_text = line if line.startswith("$") else f"${line}$"
        ax.text(0.5, y_pos, formula_text, ha="center", va="center", fontsize=20)

    fig.savefig(output_path, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)
    return output_path


def _render_cases_formula(spec: dict, output_path: Path) -> Path:
    lines = spec["lines"]
    fig_height = 0.72 * len(lines) + 0.45
    fig = plt.figure(figsize=(9.5, fig_height))
    fig.patch.set_alpha(0.0)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis("off")

    ax.text(0.18, 0.55, "$" + spec["lhs"] + "$", ha="right", va="center", fontsize=20)
    ax.text(0.22, 0.55, "{", ha="center", va="center", fontsize=58, family="DejaVu Serif")

    start_y = 0.78
    step = 0.24
    for index, line in enumerate(lines):
        ax.text(0.34, start_y - index * step, "$" + line + "$", ha="left", va="center", fontsize=19)

    fig.savefig(output_path, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)
    return output_path


def render_formula_image(spec, label: str, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{sanitize_formula_label(label)}.png"

    if isinstance(spec, dict) and spec.get("kind") == "cases":
        return _render_cases_formula(spec, output_path)

    if isinstance(spec, list):
        return _render_formula_lines(spec, output_path)

    return _render_formula_lines([spec], output_path)


def add_equation_assets(document: Document, block: dict, formula_dir: Path) -> None:
    document.add_paragraph(f"{block['label']} LaTeX / Markdown:")
    add_code_paragraph(document, f"${block['md']}$")

    image_path = render_formula_image(block["image"], block["label"], formula_dir)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(5.8))


def add_equation_paragraph(document: Document, expression: str, label: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3.2), WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("\t")

    add_markup_runs(paragraph, expression, "Cambria Math", "Cambria Math", math_mode=True)

    paragraph.add_run("\t")
    label_run = paragraph.add_run(label)
    set_run_font(label_run, "Times New Roman", "宋体")


def write_markdown(output_path: Path) -> Path:
    try:
        output_path.write_text(build_markdown_content(), encoding="utf-8")
        return output_path
    except PermissionError:
        fallback_path = output_path.with_stem(f"{output_path.stem}_updated")
        fallback_path.write_text(build_markdown_content(), encoding="utf-8")
        return fallback_path


def write_docx(output_path: Path, formula_dir: Path) -> Path:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    document.add_heading("IGWO 模型补充说明", level=0)

    for title, blocks in build_sections():
        document.add_heading(title, level=1)
        for block in blocks:
            block_type = block["type"]
            if block_type == "bullet":
                paragraph = document.add_paragraph(style="List Bullet")
                add_markup_runs(paragraph, block["docx"], "Times New Roman", "宋体")
            elif block_type == "equation":
                add_equation_assets(document, block, formula_dir)
            else:
                paragraph = document.add_paragraph()
                add_markup_runs(paragraph, block["docx"], "Times New Roman", "宋体")

    try:
        document.save(output_path)
        return output_path
    except PermissionError:
        fallback_path = output_path.with_stem(f"{output_path.stem}_updated")
        document.save(fallback_path)
        return fallback_path


def generate_modeling_documents(output_dir: Path | None = None) -> tuple[Path, Path]:
    target_dir = Path(output_dir) if output_dir is not None else Path(Config.RESULT_DIR)
    target_dir.mkdir(parents=True, exist_ok=True)

    markdown_target = target_dir / "igwo_modeling_extension.md"
    docx_target = target_dir / "igwo_modeling_extension_latex_images.docx"
    formula_dir = target_dir / "igwo_formula_images"

    markdown_path = write_markdown(markdown_target)
    docx_path = write_docx(docx_target, formula_dir)
    return markdown_path, docx_path


if __name__ == "__main__":
    md_path, docx_path = generate_modeling_documents()
    print(f"Markdown written to: {md_path}")
    print(f"DOCX written to: {docx_path}")
